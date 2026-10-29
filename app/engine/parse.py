"""Extract plain text from uploaded drafts (.txt, .md, .docx, .pdf)."""
import io

ALLOWED = {".txt", ".md", ".docx", ".pdf"}
MAX_BYTES = 2 * 1024 * 1024


class ParseError(Exception):
    pass


def extract_text(filename, data):
    name = (filename or "").lower()
    ext = name[name.rfind("."):] if "." in name else ""
    if ext not in ALLOWED:
        raise ParseError("Upload a .txt, .md, .docx or .pdf file.")
    if len(data) > MAX_BYTES:
        raise ParseError("Files must be 2 MB or smaller.")
    if ext in (".txt", ".md"):
        for enc in ("utf-8", "utf-16", "latin-1"):
            try:
                return data.decode(enc).strip()
            except UnicodeDecodeError:
                continue
    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(p for p in parts).strip()
        except Exception as exc:
            raise ParseError(f"Could not read this Word file ({type(exc).__name__}).")
    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages[:30]]
            text = "\n".join(pages).strip()
            if not text:
                raise ParseError("This PDF has no selectable text (it may be a scan).")
            return text
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError(f"Could not read this PDF ({type(exc).__name__}).")
    raise ParseError("Unsupported file.")
